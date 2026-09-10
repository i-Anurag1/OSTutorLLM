from pathlib import Path
import hashlib,json,re,math
import numpy as np
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity
DATA=Path('/app/data') if Path('/app/data').exists() else Path(__file__).resolve().parents[1]/'data'
DOCS=DATA/'documents'; INDEX=DATA/'index'; DOCS.mkdir(parents=True,exist_ok=True); INDEX.mkdir(parents=True,exist_ok=True)
vectorizer=None; matrix=None; metadata=[]
def extract(path):
    s=path.suffix.lower()
    if s in {'.txt','.md'}: return path.read_text(encoding='utf-8',errors='ignore')
    if s=='.pdf':
        from pypdf import PdfReader
        return '\n'.join((p.extract_text() or '') for p in PdfReader(str(path)).pages)
    if s=='.docx':
        from docx import Document
        return '\n'.join(p.text for p in Document(str(path)).paragraphs)
    if s=='.pptx':
        from pptx import Presentation
        return '\n'.join('\n'.join(sh.text for sh in sl.shapes if hasattr(sh,'text')) for sl in Presentation(str(path)).slides)
    return ''
def chunks(text,size=1100,overlap=180):
    clean=re.sub(r'\s+',' ',text).strip()
    if not clean: return []
    return [clean[i:i+size] for i in range(0,len(clean),size-overlap)]
def ingest_all():
    global vectorizer,matrix,metadata
    rows=[]
    for p in sorted(DOCS.rglob('*')):
        if not p.is_file(): continue
        try: raw=p.read_bytes(); text=extract(p)
        except Exception: continue
        checksum=hashlib.sha256(raw).hexdigest()
        for i,c in enumerate(chunks(text)):
            rows.append({'text':c,'source':p.name,'path':str(p.relative_to(DOCS)),'chunk':i,'checksum':checksum,'type':p.suffix.lower()[1:]})
    if not rows:
        metadata=[]; vectorizer=None; matrix=None; return {'chunks':0,'sources':0}
    texts=[r['text'] for r in rows]
    vectorizer=TfidfVectorizer(stop_words='english',ngram_range=(1,2),sublinear_tf=True)
    matrix=vectorizer.fit_transform(texts)
    metadata=rows
    np.savez_compressed(INDEX/'matrix.npz',data=matrix.toarray())
    (INDEX/'vectorizer.json').write_text(json.dumps({'vocabulary':vectorizer.vocabulary_,'idf':vectorizer.idf_.tolist()}))
    (INDEX/'metadata.json').write_text(json.dumps(rows,indent=2))
    return {'chunks':len(rows),'sources':len({r['path'] for r in rows})}
def _ensure():
    global vectorizer,matrix,metadata
    if metadata: return
    mp=INDEX/'metadata.json'; vp=INDEX/'vectorizer.json'; xp=INDEX/'matrix.npz'
    if not (mp.exists() and vp.exists() and xp.exists()): ingest_all(); return
    metadata=json.loads(mp.read_text()); cfg=json.loads(vp.read_text()); vectorizer=TfidfVectorizer(vocabulary=cfg['vocabulary']); vectorizer.idf_=np.array(cfg['idf']); vectorizer._tfidf.idf_=vectorizer.idf_; vectorizer.fixed_vocabulary_=True; matrix=np.load(xp)['data']
def _keyword_score(q,text):
    terms=set(re.findall(r'[a-zA-Z0-9]+',q.lower())); words=set(re.findall(r'[a-zA-Z0-9]+',text.lower())); return len(terms&words)/max(1,len(terms))
def retrieve(query,k=8,doc_type=None):
    _ensure()
    if not metadata: return []
    candidate_idx=[i for i,m in enumerate(metadata) if not doc_type or m['type']==doc_type]
    qv=vectorizer.transform([query]); sims=cosine_similarity(qv,matrix[candidate_idx]).ravel() if candidate_idx else np.array([])
    rows=[]
    for pos,idx in enumerate(candidate_idx):
        r=metadata[idx].copy(); r['vector_score']=float(sims[pos]); r['keyword_score']=_keyword_score(query,r['text']); r['score']=0.78*r['vector_score']+0.22*r['keyword_score']; rows.append(r)
    rows.sort(key=lambda x:x['score'],reverse=True)
    return rows[:k]
