import json
import os
import shutil
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
_TMP = tempfile.mkdtemp(prefix="ostutor-test-")
shutil.copytree(ROOT / "data" / "documents", Path(_TMP) / "documents")
os.environ["OSTUTOR_DATA_DIR"] = _TMP
os.environ["OSTUTOR_TEST_MODE"] = "1"

import numpy as np  # noqa: E402

from app import evaluation, rag  # noqa: E402
from app.grounding import extractive_answer, grounded_validate, strip_unknown_citations, support_report  # noqa: E402
from app.textproc import expand_query, sentences, stem, tokenize  # noqa: E402
from app.tutor import abstain_message, build_messages, contextualize, suggestions  # noqa: E402

_ORIG = (rag.DATA, rag.DOCS, rag.INDEX)


def setUpModule():
    """Point the retrieval engine at a throw-away copy even if another test module imported it first."""
    rag.DATA, rag.DOCS, rag.INDEX = Path(_TMP), Path(_TMP) / "documents", Path(_TMP) / "index"
    rag.INDEX.mkdir(exist_ok=True)
    rag.reset()


def tearDownModule():
    rag.DATA, rag.DOCS, rag.INDEX = _ORIG
    rag.reset()
    shutil.rmtree(_TMP, ignore_errors=True)


BENCH = json.loads((ROOT / "data" / "benchmarks" / "os_questions.json").read_text("utf-8"))


def hashing_embedder(texts, is_query):
    """Deterministic bag-of-words embedder so the dense path can be tested without downloading a model."""
    out = np.zeros((len(texts), 256), dtype="float32")
    for i, t in enumerate(texts):
        for tok in tokenize(t):
            out[i, hash(tok) % 256] += 1.0
    return out


class TextProc(unittest.TestCase):
    def test_stemmer_families(self):
        for a, b in [("paging", "pages"), ("scheduler", "scheduling"), ("buffer", "buffering"), ("frames", "framing")]:
            self.assertEqual(stem(a), stem(b))

    def test_expansion(self):
        _, added = expand_query("How does LRU use the TLB?")
        self.assertIn("least recently used", added)
        self.assertIn("translation lookaside buffer", added)

    def test_sentences_keep_decimals(self):
        self.assertEqual(len(sentences("The average is 5.67 ms. It is fine, e.g. here.")), 2)


class Chunking(unittest.TestCase):
    def test_sections_and_titles(self):
        rows, skipped = rag.build_rows()
        self.assertEqual(skipped, [])
        self.assertGreater(len(rows), 80)
        self.assertTrue(all(r["title"] and r["section"] for r in rows))
        sched = [r for r in rows if r["source"] == "02_process_scheduling.md"]
        self.assertIn("Round Robin (RR)", {r["section"] for r in sched})
        self.assertEqual(sched[0]["title"], "Process Management and CPU Scheduling")

    def test_chunk_size_and_clean_overlap(self):
        body = "\n\n".join(f"Paragraph {i}. " + "word " * 120 for i in range(6))
        chunks = rag.pack(body, 900, 120)
        self.assertGreater(len(chunks), 3)
        self.assertTrue(all(len(c) <= 900 + 125 for c in chunks))
        self.assertTrue(all(not c.startswith(("ord", "rd ")) for c in chunks))

    def test_single_long_sentence_is_split(self):
        self.assertTrue(all(len(c) <= 300 for c in rag.pack("a" * 5 + " " + ("longword " * 200), 300, 0)))

    def test_code_fences_hide_hash_comments(self):
        secs = rag._sections("# T\n\n## A\n```\n# not a heading\n```\ntext\n\n## B\nmore", True)
        self.assertEqual([s[0][-1] for s in secs], ["A", "B"])


class Retrieval(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        rag.reset()
        rag.set_embedder(None)
        os.environ["OSTUTOR_TEST_MODE"] = "1"
        rag.ingest_all()

    def top_sources(self, q, k=3):
        return [h["source"] for h in rag.search(q, k=k)["hits"]]

    def test_spot_queries(self):
        cases = {"What is the convoy effect?": "02_process_scheduling.md", "Explain Belady's anomaly": "07_page_replacement.md",
                 "how does a semaphore work": "04_synchronization.md", "what is an inode": "09_file_systems.md",
                 "difference between SCAN and C-SCAN": "10_disk_scheduling.md", "chmod 755": "11_linux_shell.md",
                 "what is a TLB": "06_memory_management.md", "banker's algorithm safe sequence": "05_deadlocks.md"}
        for q, src in cases.items():
            self.assertIn(src, self.top_sources(q), q)

    def test_hits_carry_sections_and_scores(self):
        h = rag.search("round robin quantum", k=3)["hits"][0]
        self.assertTrue(h["section"])
        self.assertIn("bm25", h["scores"])
        self.assertEqual(h["rank"], 1)

    def test_abstains_off_topic_not_on_topic(self):
        self.assertTrue(rag.search("Who won the football world cup?")["meta"]["abstain"])
        self.assertTrue(rag.search("recipe for sourdough bread")["meta"]["abstain"])
        self.assertFalse(rag.search("What is a page fault?")["meta"]["abstain"])

    def test_filters_and_empty_query(self):
        hits = rag.search("scheduling", k=5, filters={"source": "10_disk_scheduling.md"})["hits"]
        self.assertTrue(hits and all(h["source"] == "10_disk_scheduling.md" for h in hits))
        self.assertEqual(rag.search("   ")["hits"], [])
        self.assertEqual(rag.search("x", filters={"source": "nope.md"})["hits"], [])

    def test_mmr_diversifies(self):
        srcs = [h["section"] for h in rag.search("page replacement algorithms", k=5)["hits"]]
        self.assertEqual(len(srcs), len(set(srcs)))

    def test_benchmark_regression_guard(self):
        r = evaluation.evaluate_retrieval(BENCH, lambda q, k: rag.search(q, k=k))["aggregate"]
        self.assertGreaterEqual(r["recall_at_k"], 0.95)
        self.assertGreaterEqual(r["mrr"], 0.9)
        self.assertEqual(r["false_abstain_rate"], 0.0)
        self.assertGreaterEqual(r["abstention_accuracy"], 0.6)


class DenseAndPersistence(unittest.TestCase):
    def setUp(self):
        os.environ.pop("OSTUTOR_TEST_MODE", None)
        rag.reset()
        rag.set_embedder(hashing_embedder)

    def tearDown(self):
        os.environ["OSTUTOR_TEST_MODE"] = "1"
        rag.set_embedder(None)
        rag.reset()

    def test_hybrid_fusion_and_persistence(self):
        summary = rag.ingest_all()
        self.assertTrue(summary["dense_ready"])
        res = rag.search("what is the convoy effect", k=4, use_reranker=False)
        self.assertEqual(res["meta"]["mode"], "hybrid")
        self.assertTrue(res["meta"]["dense_used"])
        self.assertIn("02_process_scheduling.md", [h["source"] for h in res["hits"]])
        self.assertIsNotNone(res["hits"][0]["scores"]["dense"])
        rag.reset()
        self.assertTrue(rag.index_status()["dense_ready"])  # reloaded from disk with embeddings

    def test_missing_embedder_degrades_to_sparse(self):
        rag.set_embedder(lambda t, q: (_ for _ in ()).throw(RuntimeError("no model")))
        summary = rag.ingest_all()
        self.assertFalse(summary["dense_ready"])
        res = rag.search("what is an inode", k=3)
        self.assertEqual(res["meta"]["mode"], "sparse")
        self.assertTrue(res["hits"])

    def test_fingerprint_change_triggers_rebuild_and_stale_memory_reloads(self):
        rag.ingest_all()
        n_before = rag.index_status()["chunks"]
        stale, stale_mtime = rag._STATE, rag._SEEN_MTIME
        (rag.DOCS / "99_extra.md").write_text("# Extra\n\n## Zebra scheduling\nThe zebra scheduler assigns stripes to processes.\n", encoding="utf-8")
        try:
            status = rag.ensure_ready()  # fingerprint differs -> rebuild
            self.assertGreater(status["chunks"], n_before)
            # simulate another process (the worker) having re-indexed while this one still holds old state
            rag._STATE, rag._SEEN_MTIME = stale, stale_mtime
            hits = rag.search("zebra scheduler stripes", k=3)["hits"]
            self.assertEqual(hits[0]["source"], "99_extra.md")
        finally:
            (rag.DOCS / "99_extra.md").unlink()
            rag.ingest_all()

    def test_corrupt_manifest_is_ignored(self):
        rag.ingest_all()
        (rag.INDEX / "manifest.json").write_text("{not json", encoding="utf-8")
        rag.reset()
        self.assertTrue(rag.search("what is a thread", k=2)["hits"])  # rebuilt transparently


class Extraction(unittest.TestCase):
    def test_docx_and_pptx_headings_become_sections(self):
        from docx import Document
        from pptx import Presentation

        d = Path(tempfile.mkdtemp())
        doc = Document()
        doc.add_heading("Lab Manual", 1)
        doc.add_heading("Semaphores in practice", 2)
        doc.add_paragraph("Use wait and signal to guard a critical section.")
        doc.save(d / "lab.docx")
        rows = []
        for page in rag._extract_pages(d / "lab.docx"):
            rows += rag._sections(page["text"], page["markdown"])
        self.assertEqual(rows[0][0], ["Lab Manual", "Semaphores in practice"])
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = "Deadlock slide"
        slide.placeholders[1].text = "Four Coffman conditions"
        prs.save(d / "s.pptx")
        text = rag.extract(d / "s.pptx")
        self.assertIn("# Deadlock slide", text)
        self.assertIn("Coffman", text)


class Grounding(unittest.TestCase):
    EV = [{"text": "FCFS runs processes in arrival order and is non-preemptive. It suffers from the convoy effect.", "section": "FCFS"},
          {"text": "Round Robin gives each process a fixed time quantum before moving it to the back of the queue.", "section": "Round Robin"}]

    def test_support_flags_unsupported_claims(self):
        a = "FCFS runs processes in arrival order [S1]. Bananas are yellow [S1]. Round Robin uses a time quantum [S2]."
        flags = [s["supported"] for s in support_report(a, self.EV)["sentences"]]
        self.assertEqual(flags, [True, False, True])
        self.assertFalse(grounded_validate(a, [{"id": "S1"}, {"id": "S2"}], self.EV))

    def test_good_answer_validates(self):
        a = "FCFS runs processes in arrival order [S1]. Round Robin gives each process a time quantum [S2]."
        self.assertTrue(grounded_validate(a, [{"id": "S1"}, {"id": "S2"}], self.EV))

    def test_unknown_citation_rejected_and_stripped(self):
        self.assertFalse(grounded_validate("x y z [S9]", [{"id": "S1"}], self.EV))
        self.assertEqual(strip_unknown_citations("Hello [S1] there [S4].", 2), "Hello [S1] there.")

    def test_extractive_answer_is_cited_and_relevant(self):
        text, used = extractive_answer("what is the convoy effect", [{"text": self.EV[0]["text"]}, {"text": self.EV[1]["text"]}])
        self.assertIn("convoy", text)
        self.assertIn("[S1]", text)
        self.assertEqual(extractive_answer("x", []), ("", []))


class Tutor(unittest.TestCase):
    def test_contextualize_followups(self):
        hist = [{"role": "user", "content": "Explain FIFO page replacement"}, {"role": "assistant", "content": "..."}]
        self.assertIn("FIFO", contextualize("what about LRU?", hist))
        self.assertEqual(contextualize("Explain the Banker's algorithm in detail please", hist), "Explain the Banker's algorithm in detail please")
        self.assertEqual(contextualize("hello there", []), "hello there")

    def test_messages_include_evidence_modes_and_history(self):
        hits = [{"source": "a.md", "section": "Sec", "text": "Evidence text."}]
        m = build_messages("Q?", hits, "socratic", "hard", "Paging", 80, [{"role": "user", "content": "earlier"}])
        self.assertEqual(m[0]["role"], "system")
        self.assertIn("Socratic", m[0]["content"])
        self.assertIn("advanced", m[0]["content"])
        self.assertIn("[S1] (a.md › Sec)", m[-1]["content"])
        self.assertEqual(m[1]["content"], "earlier")

    def test_adaptive_level_follows_mastery(self):
        self.assertIn("beginner", build_messages("q", [], "teach", "adaptive", None, 10)[0]["content"])
        self.assertIn("advanced", build_messages("q", [], "teach", "adaptive", None, 90)[0]["content"])

    def test_suggestions_and_abstain(self):
        s = suggestions([{"section": "LRU page replacement"}, {"section": "Clock and second chance"}], "explain lru")
        self.assertEqual(s, ["Explain clock and second chance"])
        self.assertIn("will not guess", abstain_message("x"))


if __name__ == "__main__":
    unittest.main()
